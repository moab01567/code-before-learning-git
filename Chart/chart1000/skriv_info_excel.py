import openpyxl

class Excel:
	def __init__(self):
		self.info = self.fiks_data()
		self.sheet = self.fiks_excelfil()
		self.hoved()



	def hoved(self):
		self.skriv_rad = 1
		for day in self.info:
			antall = len(day)
			self.skriv_excel(self.sheet, day)
			self.skriv_rad = self.skriv_rad + 1
		self.wb.save("nov2021.xlsx")


	def skriv_excel(self, sheet, trades):
		gevist = 0
		for inn, type_, s, v, pris, sw, tp, ut2, ut, sw, profitt in trades:
		
			sheet.cell(row= self.skriv_rad, column = 1).value = inn
			sheet.cell(row= self.skriv_rad, column = 2).value = type_
			sheet.cell(row= self.skriv_rad, column = 3).value = s
			sheet.cell(row= self.skriv_rad, column = 4).value = v
			sheet.cell(row= self.skriv_rad, column = 5).value = pris
			sheet.cell(row= self.skriv_rad, column = 6).value = sw
			sheet.cell(row= self.skriv_rad, column = 7).value = tp
			sheet.cell(row= self.skriv_rad, column = 8).value = ut2
			sheet.cell(row= self.skriv_rad, column = 9).value = ut
			sheet.cell(row= self.skriv_rad, column = 10).value = sw
			sheet.cell(row= self.skriv_rad, column = 11).value = profitt
			gevist = gevist + float(profitt)
			self.skriv_rad = self.skriv_rad + 1
		sheet.cell(row= self.skriv_rad, column = 12).value = f"{gevist}"


	def fiks_excelfil(self):
		self.wb = openpyxl.Workbook()
		sheet = self.wb.active
		return sheet
	
	def fiks_data(self):
		with open("data.txt","r") as r:
			info = r.read()

		info = info.replace("'","")


		res = info.strip('][').split(', ')


		print(res)
		mellomliste = []
		antall = 0
		res2 = []
		for i in res:
			mellomliste.append(i)
			antall = antall + 1
			
			if antall == 11:
				res2.append(mellomliste)
				antall = 0
				mellomliste = []
			
			else:
				pass



		all_trades = []
		day = []
		one_trade = []
		tid = res2[0][0][0:10]
		for inn, type_, s, v, pris, sw, tp, ut2, ut, sw, profitt in res2:
			if inn[0:10] == tid:
				one_trade.append(inn)
				one_trade.append(type_)
				one_trade.append(s)
				one_trade.append(v)
				one_trade.append(pris)
				one_trade.append(sw)
				one_trade.append(tp)
				one_trade.append(ut2)
				one_trade.append(ut)
				one_trade.append(sw)
				one_trade.append(profitt)
				
				day.append(one_trade)
				one_trade = []

			else:
				all_trades.append(day)
				day = []
				tid = inn[0:10]

				one_trade.append(inn)
				one_trade.append(type_)
				one_trade.append(s)
				one_trade.append(v)
				one_trade.append(pris)
				one_trade.append(sw)
				one_trade.append(tp)
				one_trade.append(ut2)
				one_trade.append(ut)
				one_trade.append(sw)
				one_trade.append(profitt)
				
				day.append(one_trade)
				one_trade = []

			print(inn, type_, s, v, pris,sw,tp , ut2, ut, sw, profitt)

		all_trades.append(day)


		print()
		print()
		print()
		print()
		print()
		print(all_trades)
		return all_trades



start = Excel()






